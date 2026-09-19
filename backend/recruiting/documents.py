import re
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

MAX_BYTES = 5 * 1024 * 1024
MAX_TEXT = 40000


def extract_resume(upload):
    if upload.size > MAX_BYTES:
        raise ValueError("File exceeds 5 MB.")
    suffix = Path(upload.name).suffix.lower()
    try:
        if suffix == ".pdf":
            reader = PdfReader(upload)
            if reader.is_encrypted or len(reader.pages) > 50:
                raise ValueError("Use an unencrypted PDF with at most 50 pages.")
            parts = []
            for page in reader.pages:
                if (
                    len(page.get_contents().get_data() if page.get_contents() else b"")
                    > 5 * MAX_BYTES
                ):
                    raise ValueError("PDF page is too complex. Please upload plain text.")
                parts.append(page.extract_text() or "")
                if sum(map(len, parts)) > MAX_TEXT:
                    raise ValueError("Resume exceeds 40,000 characters.")
            text = "\n\n".join(parts)
        elif suffix == ".docx":
            with ZipFile(upload) as archive:
                if sum(item.file_size for item in archive.infolist()) > 20 * MAX_BYTES:
                    raise ValueError("DOCX expands beyond the supported size.")
            upload.seek(0)
            doc = Document(upload)
            text = "\n".join(
                [p.text for p in doc.paragraphs]
                + [
                    " | ".join(c.text for c in row.cells)
                    for table in doc.tables
                    for row in table.rows
                ]
            )
        elif suffix == ".txt":
            text = upload.read().decode("utf-8-sig")
        else:
            raise ValueError("Use PDF, DOCX, or UTF-8 TXT files.")
    except (ValueError, UnicodeDecodeError, BadZipFile) as exc:
        raise ValueError(str(exc)) from exc
    except Exception as exc:
        raise ValueError(
            "Could not read this document. Try a text-based PDF or paste its text."
        ) from exc
    text = text.replace("\x00", "").strip()
    if len(text) < 40:
        raise ValueError(
            "Not enough readable text. Scanned PDFs need OCR; paste the resume text instead."
        )
    if len(text) > MAX_TEXT:
        raise ValueError("Resume exceeds 40,000 characters.")
    upload.seek(0)
    return text


def contact_from_text(text, filename):
    email = re.search(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", text, re.I)
    first = text.splitlines()[0].strip()
    name = (
        first
        if 2 <= len(first) <= 100 and "@" not in first
        else Path(filename).stem.replace("_", " ")
    )
    return name[:160], email.group(0)[:254] if email else ""
